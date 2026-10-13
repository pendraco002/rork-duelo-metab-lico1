"""Monta o PPTX a partir das imagens dos slides (geradas pelo build.mjs).

Cada slide é uma imagem em tela cheia + anotações do apresentador (roteiro de fala).
"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
TMP = os.path.join(HERE, ".tmp")
OUT = os.path.join(os.path.dirname(HERE), "Apresentacao_Whey_Protein_Concentrado.pptx")

NOTES = {
    1: "[Bruna] Bom dia. Somos o Grupo 1 e vamos falar do whey protein concentrado. O gráfico de pontos mostra o que tem em 100 g de WPC80: cada ponto é 1 grama — 80 são proteína.",
    2: "[Bruna] Organizamos a apresentação em seis perguntas, uma para cada tópico pedido. Eu começo com o que é o whey e quem consome.",
    3: "[Bruna] O whey nasce de um resíduo. De cada 10 litros de leite que viram queijo, cerca de 9 viram soro. Esse soro leva cerca de 20% das proteínas do leite e é muito poluente se for descartado: DBO de 30 a 60 mil mg/L.",
    4: "[Bruna] WPC é o soro sem boa parte da água, da lactose e dos minerais, com 25 a 89% de proteína em base seca. O dos suplementos é o WPC80. A matéria-prima é o soro doce, de queijos de coalho. No pote entram também cacau ou aromas, edulcorantes, gomas e lecitina. O isolado (WPI) e o hidrolisado (WPH) são variações.",
    5: "[Bruna] O mesmo produto tem quatro leituras: para a ANVISA é suplemento alimentar e, desde 1º de setembro de 2026, precisa de notificação; para o MAPA é ingrediente lácteo; tecnologicamente é um pó feito por membranas e secagem; pela classificação NOVA, o saborizado é ultraprocessado.",
    6: "[Bruna] Consumo: em 2020, 59% dos lares tinham alguém usando suplemento (dado de suplementos em geral, não só whey). Os concentrados de proteína cresceram 3% em 2024. Nas academias, a proteína está entre os mais usados, muitas vezes sem orientação de nutricionista. Passo para o Thiago.",
    7: "[Thiago] O WPC80 tem 80 a 82% de proteína, mas ainda tem lactose: de 4 a 10%. Metade da proteína é beta-lactoglobulina. Por que uma dose de 30 g dá cerca de 23 g de proteína? Porque os 80% são em base seca e o pó saborizado tem umidade e ingredientes de sabor. A qualidade proteica é máxima: PDCAAS 1,0, rica em leucina.",
    8: "[Thiago] Este é o esquema do rótulo que analisamos, com os seis elementos que a lei exige: denominação, tabela nutricional, ingredientes, alergênicos e lactose, as três advertências da RDC 243 e o número de notificação na ANVISA.",
    9: "[Thiago] Separando o que o rótulo diz do que nós calculamos: 76,7% do pó é proteína; 75% das calorias vêm da proteína; ele não tem lupa porque fica abaixo dos limites; e os 23 g por dose estão dentro da faixa de 20 a 40 g. Passo para o Gabriel.",
    10: "[Gabriel] O processo tem duas fábricas. A linha azul é o laticínio: recepção e resfriamento do soro, clarificação, desnate, pasteurização (que é PCC), resfriamento, ultrafiltração e diafiltração, concentração, spray drying, aglomeração e detector de metais (outro PCC). O WPC80 em pó vira insumo da segunda fábrica, a de suplementos (linha preta): recepção, pesagem, mistura, detector de metais, envase, rotulagem e estoque.",
    11: "[Gabriel] Três etapas definem o WPC. A pasteurização, a 72–75 °C por 15–20 s, elimina patógenos com pouca desnaturação. A ultrafiltração segura a proteína e deixa passar lactose e sais; a diafiltração lava o retido com água. No spray drying, o ar entra a 175–200 °C, mas a gota fica fria enquanto evapora. Passo para o Renan.",
    12: "[Renan] Esta matriz mostra o que cada etapa muda: efeitos físicos, químicos, microbiológicos, sensoriais e nutricionais. Os dois pontos destacados são a reação de Maillard e a perda de lisina disponível, que acontecem principalmente na estocagem. E tudo depende de temperatura, tempo, umidade, oxigênio e formulação.",
    13: "[Renan] Um ponto importante para não generalizar: desnaturar não é estragar. O calor muda a forma da proteína, mas os aminoácidos continuam lá. A perda real vem da reação de Maillard, quando a lactose se liga à lisina e o corpo deixa de aproveitá-la.",
    14: "[Renan] Perigos biológicos, químicos e físicos e seus controles. Os PCCs são a pasteurização e o detector de metais. A base são as boas práticas e os POPs. E um alerta: Salmonella e Cronobacter não crescem no pó, mas sobrevivem por meses.",
    15: "[Renan] Conservação: o inimigo do pó é a água. Quando o pó ganha umidade ou calor, passa da transição vítrea, fica pegajoso, empedra e a Maillard acelera. Por isso: embalagem com barreira, pote bem fechado, dosador seco e local fresco. Passo para o Pedro.",
    16: "[Pedro] O rótulo nem sempre conta a verdade. Um estudo da UFMG achou erro de rótulo em 88% das marcas e uma marca com proteína 35% abaixo da tolerância. Em Recife, só 2 de 14 amostras estavam totalmente conformes. Em 2025, a ANVISA suspendeu um lote com glúten não declarado. E existe a fraude do amino spiking, que infla a proteína no laudo.",
    17: "[Pedro] Vale a pena? Depende. A favor: proteína de alta qualidade, praticidade, ajuda quando a comida não basta e aproveita um resíduo. Contra: não é indispensável, o saborizado é ultraprocessado, tem lactose e é alergênico, e o mercado tem falhas. O contexto importa: frequência, dose e quem vai usar.",
    18: "[Pedro] Nossa conclusão: nem vilão, nem indispensável. O whey pode entrar na alimentação cotidiana quando complementa as refeições, quando existe necessidade real avaliada por nutricionista, e quando o produto é notificado e o rótulo foi conferido.",
    19: "Referências — legislação.",
    20: "Referências — artigos, livros e dados.",
    21: "[Todos] Agradecer e abrir para perguntas. Respostas prontas na seção 7 do dossiê.",
}

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
blank = prs.slide_layouts[6]

pngs = sorted(f for f in os.listdir(os.path.join(TMP, "slides")) if f.endswith(".png"))
for i, name in enumerate(pngs, start=1):
    s = prs.slides.add_slide(blank)
    jpg = os.path.join(TMP, "slides", name[:-4] + ".jpg")
    Image.open(os.path.join(TMP, "slides", name)).convert("RGB").resize((1920, 1080), Image.LANCZOS).save(jpg, quality=90, optimize=True)
    s.shapes.add_picture(jpg, 0, 0, prs.slide_width, prs.slide_height)
    s.notes_slide.notes_text_frame.text = NOTES.get(i, "")

prs.save(OUT)
print("pptx:", OUT, len(pngs), "slides")
